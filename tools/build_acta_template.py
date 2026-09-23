from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("dist/modelo-acta-asembleas.docx").resolve()
LOGO = Path(
    "/Users/a.moron/Library/CloudStorage/GoogleDrive-asembleasensino@gmail.com/"
    "Mi unidad/online/redes/AAEP Identidade/01_LOGOS-OFICIAIS/logotipo/"
    "logotipo-horizontal-cor.png"
)

BLACK = "221F1F"
RED = "D64B43"
OFFWHITE = "F7F5F1"
LIGHT_GRAY = "E8E5DF"
MID_GRAY = "6D6963"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_width(cell, width_twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, pct=5000):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(pct))
    tbl_w.set(qn("w:type"), "pct")


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def remove_style_borders(style):
    p_pr = style._element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def style_run(run, size=None, bold=False, color=BLACK, italic=False):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    if size:
        run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", style=None, size=11, bold=False, color=BLACK, italic=False, before=0, after=6):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.08
    if text:
        run = paragraph.add_run(text)
        style_run(run, size=size, bold=bold, color=color, italic=italic)
    return paragraph


def add_section_heading(doc, text):
    paragraph = add_para(doc, text, style="Heading 1", size=15, bold=True, before=12, after=6)
    keep_with_next(paragraph)
    return paragraph


def add_subheading(doc, text):
    paragraph = add_para(doc, text, style="Heading 2", size=12.5, bold=True, before=8, after=4)
    keep_with_next(paragraph)
    return paragraph


def format_table(table, header_rows=1, widths=None, header_fill=BLACK):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table)
    for row_idx, row in enumerate(table.rows):
        if row_idx < header_rows:
            repeat_table_header(row)
        for col_idx, cell in enumerate(row.cells):
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                set_cell_width(cell, widths[col_idx])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    style_run(run, size=10.2, color=BLACK)
            if row_idx < header_rows:
                set_cell_shading(cell, header_fill)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        style_run(run, size=9.8, bold=True, color="FFFFFF")
            elif row_idx % 2 == 0:
                set_cell_shading(cell, "FAFAF8")


def cell_text(cell, text, bold=False, color=BLACK, size=10.2):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    style_run(run, size=size, bold=bold, color=color)


def add_label_value_table(doc, pairs):
    table = doc.add_table(rows=0, cols=4)
    widths = [1750, 2950, 1750, 2950]
    for left_label, left_value, right_label, right_value in pairs:
        row = table.add_row()
        for idx, text in enumerate((left_label, left_value, right_label, right_value)):
            cell_text(row.cells[idx], text, bold=idx in (0, 2), color=BLACK if idx in (0, 2) else MID_GRAY)
            set_cell_shading(row.cells[idx], OFFWHITE if idx in (0, 2) else "FFFFFF")
    format_table(table, header_rows=0, widths=widths)
    return table


def add_blank_rows_table(doc, headers, rows, widths, header_fill=BLACK):
    table = doc.add_table(rows=1, cols=len(headers))
    for idx, header in enumerate(headers):
        cell_text(table.rows[0].cells[idx], header, bold=True, color="FFFFFF", size=9.8)
    for _ in range(rows):
        row = table.add_row()
        for cell in row.cells:
            cell_text(cell, " ", color=MID_GRAY)
    format_table(table, widths=widths, header_fill=header_fill)
    return table


def add_notes_area(doc, lines=5):
    table = doc.add_table(rows=lines, cols=1)
    for row in table.rows:
        cell_text(row.cells[0], " ", color=MID_GRAY)
    format_table(table, header_rows=0, widths=[9400])
    return table


def set_document_defaults(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    remove_style_borders(title)

    for name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11.5)):
        style = styles[name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
        remove_style_borders(style)


def add_footer(section):
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    run = footer.add_run("Asembleas Abertas do Ensino Público · Modelo de acta")
    style_run(run, size=8.5, color=MID_GRAY)


def build_doc():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    set_document_defaults(doc)
    add_footer(section)

    logo_para = doc.add_paragraph()
    logo_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    logo_para.paragraph_format.space_after = Pt(8)
    logo_run = logo_para.add_run()
    logo_run.add_picture(str(LOGO), width=Inches(4.35))

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(3)
    title.add_run("Modelo de acta de asemblea")
    remove_paragraph_borders(title)

    subtitle = add_para(
        doc,
        "Plantilla común para asembleas comarcais, de centro e xerais",
        size=12.5,
        color=MID_GRAY,
        after=8,
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT

    intro = add_para(
        doc,
        "Este documento serve para recoller de maneira homoxénea os datos da convocatoria, "
        "as persoas asistentes, o desenvolvemento dos puntos tratados, os acordos adoptados "
        "e as tarefas pendentes. Substitúe os campos entre corchetes polo contido de cada reunión.",
        size=10.8,
        color=BLACK,
        after=10,
    )

    add_section_heading(doc, "Datos da asemblea")
    add_label_value_table(
        doc,
        [
            ("Tipo de asemblea", "☐ Comarcal   ☐ De centro   ☐ Xeral   ☐ Outra", "Data", "[dd/mm/aaaa]"),
            ("Centro ou comarca", "[Nome do centro, comarca ou ámbito]", "Lugar ou modalidade", "[Presencial / en liña / mixta]"),
            ("Hora de inicio", "[hh:mm]", "Hora de remate", "[hh:mm]"),
            ("Convoca", "[Órgano ou persoa que convoca]", "Modera", "[Nome e apelidos]"),
            ("Secretaría da acta", "[Nome e apelidos]", "Anexos", "☐ Si   ☐ Non"),
        ],
    )

    add_subheading(doc, "Persoas asistentes")
    add_blank_rows_table(
        doc,
        ["Nome e apelidos", "Centro ou ámbito", "Responsabilidade ou representación", "Sinatura"],
        7,
        [2600, 2200, 2750, 1850],
        header_fill=BLACK,
    )

    add_para(doc, "Persoas que xustifican ausencia", size=10.5, bold=True, before=6, after=3)
    add_notes_area(doc, lines=3)

    add_section_heading(doc, "Orde do día")
    add_blank_rows_table(
        doc,
        ["Nº", "Punto previsto", "Documentación ou referencia"],
        6,
        [700, 5700, 3000],
        header_fill=RED,
    )

    add_section_heading(doc, "Desenvolvemento da reunión")
    add_para(
        doc,
        "Para cada punto da orde do día, recolle as intervencións principais, propostas presentadas, "
        "matices relevantes e documentación empregada. Engade ou elimina bloques segundo sexa necesario.",
        size=10.5,
        color=MID_GRAY,
        after=8,
    )

    for number in range(1, 5):
        add_subheading(doc, f"Punto {number}  [Título do punto]")
        add_notes_area(doc, lines=5 if number < 4 else 4)

    add_section_heading(doc, "Acordos adoptados")
    add_blank_rows_table(
        doc,
        ["Nº", "Acordo", "Responsable", "Prazo", "Seguimento"],
        6,
        [650, 4450, 1750, 1200, 1350],
        header_fill=BLACK,
    )

    add_section_heading(doc, "Votacións")
    add_blank_rows_table(
        doc,
        ["Proposta sometida a votación", "A favor", "En contra", "Abstencións", "Resultado"],
        4,
        [4200, 1100, 1100, 1250, 1750],
        header_fill=RED,
    )

    add_section_heading(doc, "Tarefas pendentes")
    add_blank_rows_table(
        doc,
        ["Tarefa", "Persoa ou grupo responsable", "Prazo", "Estado"],
        6,
        [4100, 2550, 1350, 1400],
        header_fill=BLACK,
    )

    add_section_heading(doc, "Rogos e preguntas")
    add_notes_area(doc, lines=4)

    add_section_heading(doc, "Peche da acta")
    add_label_value_table(
        doc,
        [
            ("Data de redacción", "[dd/mm/aaaa]", "Data de aprobación", "[dd/mm/aaaa]"),
            ("Acta aprobada por", "☐ Asemblea   ☐ Coordinación   ☐ Outro", "Número de versión", "[v1 / v2 / definitiva]"),
        ],
    )

    add_subheading(doc, "Sinaturas")
    table = doc.add_table(rows=2, cols=2)
    cell_text(table.rows[0].cells[0], "Secretaría da acta", bold=True, color="FFFFFF")
    cell_text(table.rows[0].cells[1], "Visto e prace", bold=True, color="FFFFFF")
    cell_text(table.rows[1].cells[0], "\n\n\nNome: [Nome e apelidos]\nData: [dd/mm/aaaa]", color=MID_GRAY)
    cell_text(table.rows[1].cells[1], "\n\n\nNome: [Nome e apelidos]\nCargo ou función: [Cargo]\nData: [dd/mm/aaaa]", color=MID_GRAY)
    format_table(table, widths=[4700, 4700], header_fill=BLACK)

    add_para(
        doc,
        "Anexos incorporados á acta",
        style="Heading 2",
        size=12.5,
        bold=True,
        before=10,
        after=4,
    )
    add_blank_rows_table(
        doc,
        ["Nº", "Descrición do anexo", "Formato ou ligazón"],
        4,
        [700, 5750, 2950],
        header_fill=RED,
    )

    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
